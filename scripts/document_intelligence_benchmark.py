#!/usr/bin/env python3
"""Reproducible document-intelligence benchmark harness.

The harness is deliberately provider-neutral.  It exercises the canonical
artifact contract, records structural metrics and emits ``SKIPPED`` for an
optional provider that is not installed/configured.  It never chooses a
runtime default; selection remains an explicit, benchmark-reviewed policy.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
import platform
import sys
import tempfile
import time
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_FIXTURES = ROOT / "test" / "fixtures" / "document-intelligence"
DEFAULT_REPORT = ROOT / "reports" / "document-intelligence" / "benchmark.json"

# The extractor is a Python app inside the monorepo rather than an installed
# distribution when this script is invoked from the repository root.
EXTRACTOR_ROOT = ROOT / "apps" / "extractor"
if str(EXTRACTOR_ROOT) not in sys.path:
    sys.path.insert(0, str(EXTRACTOR_ROOT))


def _digest(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_xlsx(target: Path, fixture: dict[str, Any]) -> None:
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    rel_namespace = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    rows: list[str] = []
    for row_number, values in enumerate(fixture["sheets"][0]["rows"], start=1):
        cells: list[str] = []
        for column_number, value in enumerate(values, start=1):
            column = ""
            number = column_number
            while number:
                number, remainder = divmod(number - 1, 26)
                column = chr(65 + remainder) + column
            cells.append(
                f'<c r="{column}{row_number}" t="inlineStr"><is><t>{escape(str(value))}</t></is></c>'
            )
        rows.append(f'<row r="{row_number}">{"".join(cells)}</row>')
    workbook = (
        f'<workbook xmlns="{namespace}" xmlns:r="{rel_namespace}"><sheets>'
        f'<sheet name="{escape(fixture["sheets"][0]["name"])}" sheetId="1" r:id="rId1"/>'
        "</sheets></workbook>"
    )
    worksheet = f'<worksheet xmlns="{namespace}"><sheetData>{"".join(rows)}</sheetData></worksheet>'
    rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" '
        'Target="xl/workbook.xml"/></Relationships>'
    )
    workbook_rels = (
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" '
        'Target="worksheets/sheet1.xml"/></Relationships>'
    )
    content_types = (
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
        '<Default Extension="xml" ContentType="application/xml"/>'
        '<Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>'
        '<Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>'
        '</Types>'
    )
    with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        def write_entry(name: str, content: str) -> None:
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            archive.writestr(info, content)

        write_entry("[Content_Types].xml", content_types)
        write_entry("_rels/.rels", rels)
        write_entry("xl/workbook.xml", workbook)
        write_entry("xl/_rels/workbook.xml.rels", workbook_rels)
        write_entry("xl/worksheets/sheet1.xml", worksheet)


def _pdf_object(number: int, value: bytes) -> bytes:
    return f"{number} 0 obj\n".encode() + value + b"\nendobj\n"


def _write_pdf(target: Path, fixture: dict[str, Any]) -> None:
    text = str(fixture["pages"][0]).replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        _pdf_object(1, b"<< /Type /Catalog /Pages 2 0 R >>"),
        _pdf_object(2, b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>"),
        _pdf_object(
            3,
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        ),
        _pdf_object(4, b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream"),
        _pdf_object(5, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
    ]
    header = b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n"
    body = bytearray(header)
    offsets = [0]
    for obj in objects:
        offsets.append(len(body))
        body.extend(obj)
    xref_offset = len(body)
    body.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    body.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        body.extend(f"{offset:010d} 00000 n \n".encode())
    body.extend(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode()
    )
    target.write_bytes(body)


def _normalize_zip(path: Path) -> None:
    """Normalize ZIP member timestamps so generated office fixtures hash stably."""

    normalized = path.with_suffix(path.suffix + ".normalized")
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        normalized, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for name in sorted(source.namelist()):
            info = zipfile.ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            target.writestr(info, source.read(name))
    normalized.replace(path)


def _write_docx(target: Path, fixture: dict[str, Any]) -> None:
    from docx import Document

    document = Document()
    for paragraph in fixture.get("paragraphs", []):
        document.add_paragraph(paragraph["text"], style=paragraph.get("style"))
    for rows in fixture.get("tables", []):
        table = document.add_table(rows=len(rows), cols=len(rows[0]))
        for row_number, values in enumerate(rows):
            for column_number, value in enumerate(values):
                table.cell(row_number, column_number).text = str(value)
    document.save(target)
    _normalize_zip(target)


def _write_pptx(target: Path, fixture: dict[str, Any]) -> None:
    from pptx import Presentation

    presentation = Presentation()
    for slide_fixture in fixture.get("slides", []):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = slide_fixture.get("title", "")
        slide.placeholders[1].text = slide_fixture.get("body", "")
        if slide_fixture.get("notes"):
            slide.notes_slide.notes_text_frame.text = slide_fixture["notes"]
    presentation.save(target)
    _normalize_zip(target)


def _write_png(target: Path, fixture: dict[str, Any]) -> None:
    from PIL import Image, ImageDraw

    width = int(fixture.get("width", 320))
    height = int(fixture.get("height", 180))
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, width - 10, height - 10), outline="black", width=2)
    draw.text((20, 20), str(fixture.get("label", "diagram")), fill="black")
    image.save(target, format="PNG", optimize=False)


def _materialize_fixture(fixtures: Path, entry: dict[str, Any], temp_dir: Path) -> Path:
    source = fixtures / str(entry["path"])
    generated = entry.get("generated_format")
    if not generated:
        return source
    payload = json.loads(source.read_text(encoding="utf-8"))
    target = temp_dir / f"{entry['id']}.{generated}"
    if generated == "xlsx":
        _write_xlsx(target, payload)
    elif generated == "pdf":
        _write_pdf(target, payload)
    elif generated == "docx":
        _write_docx(target, payload)
    elif generated == "pptx":
        _write_pptx(target, payload)
    elif generated == "png":
        _write_png(target, payload)
    else:
        raise ValueError(f"unknown generated fixture format: {generated}")
    return target


def _full_text(artifact: Any) -> str:
    values = [artifact.text_content()]
    for table in artifact.tables:
        values.append("\n".join(" | ".join(row) for row in [table.headers, *table.rows]))
    for page in artifact.pages:
        values.append(page.text or "")
    return "\n".join(values)


def _metrics(artifact: Any, expected: dict[str, Any]) -> dict[str, Any]:
    full_text = _full_text(artifact).lower()
    snippets = [str(value).lower() for value in expected.get("text", [])]
    matched = sum(1 for snippet in snippets if snippet in full_text)
    table_rows = sum(len(table.rows) for table in artifact.tables)
    expected_table_count = expected.get("table_count")
    expected_table_rows = expected.get("table_rows")
    table_structure_match = (
        expected_table_count is None or len(artifact.tables) == expected_table_count
    ) and (expected_table_rows is None or table_rows == expected_table_rows)
    expected_heading_count = expected.get("heading_count")
    return {
        "text_recall": (matched / len(snippets)) if snippets else None,
        "heading_hierarchy": (
            1.0 if len(artifact.headings) == expected_heading_count else 0.0
        )
        if expected_heading_count is not None
        else None,
        "heading_count": len(artifact.headings),
        "table_structure_accuracy": 1.0 if table_structure_match else 0.0,
        "table_count": len(artifact.tables),
        "table_rows": table_rows,
        "formula_preservation": (
            1.0 if len(artifact.equations) == expected["equation_count"] else 0.0
        )
        if "equation_count" in expected
        else None,
        "figure_detection": (
            1.0 if len(artifact.figures) == expected["figure_count"] else 0.0
        )
        if "figure_count" in expected
        else None,
        "locator_accuracy": 1.0 if (artifact.locators and (artifact.blocks or artifact.pages)) else 0.0,
        "chunk_boundary_quality": (
            len(artifact.reading_order) / len(artifact.blocks)
            if artifact.blocks
            else 0.0
        ),
        "structured_units": len(artifact.blocks),
        "page_count": len(artifact.pages),
    }


def _versions() -> dict[str, str]:
    versions: dict[str, str] = {}
    for package in ("pydantic", "fastapi", "pypdf", "docling", "marker-pdf", "httpx"):
        try:
            versions[package] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            versions[package] = "NOT_INSTALLED"
    return versions


def _media_supported(availability: Any, media_type: str) -> bool:
    aliases = {
        "text/markdown": {"markdown-text"},
        "text/plain": {"markdown-text", "code"},
        "text/html": {"html-snapshot"},
        "application/pdf": {"pdf"},
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": {"docx"},
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": {"pptx"},
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": {"xlsx"},
    }
    candidates = set(aliases.get(media_type, set()))
    candidates.add(media_type)
    if media_type.startswith("image/"):
        candidates.add("image-metadata")
        candidates.add("image")
    return bool(set(availability.media) & candidates)


def run(fixtures: Path, report: Path, repeats: int) -> dict[str, Any]:
    manifest = json.loads((fixtures / "manifest.json").read_text(encoding="utf-8"))
    generated_at = datetime.now(UTC).isoformat()
    metadata: dict[str, Any] = {
        "schema_version": "1.0",
        "generated_at": generated_at,
        "command": " ".join(sys.argv),
        "python": platform.python_version(),
        "platform": platform.platform(),
        "versions": _versions(),
        "fixture_root": str(fixtures),
        "repeats": repeats,
    }
    try:
        from app.adapters.deterministic import DeterministicTextAdapter
        from app.adapters.optional import ChunkrAdapter, DoclingAdapter, MarkerAdapter

        adapters = [
            DeterministicTextAdapter(),
            DoclingAdapter(),
            MarkerAdapter(),
            ChunkrAdapter(),
        ]
    except Exception as error:  # noqa: BLE001  # dependency bootstrap evidence
        result = {
            "metadata": metadata,
            "status": "BLOCKED",
            "reason": f"extractor import failed: {error}",
            "fixtures": [],
            "selection": {
                "status": "NOT_SELECTED",
                "reason": "benchmark could not import extractor dependencies",
            },
        }
        report.parent.mkdir(parents=True, exist_ok=True)
        report.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        return result

    rows: list[dict[str, Any]] = []
    with tempfile.TemporaryDirectory(prefix="akp-document-intelligence-") as temp:
        temp_dir = Path(temp)
        for entry in manifest["fixtures"]:
            source = _materialize_fixture(fixtures, entry, temp_dir)
            source_hash = _digest(source)
            for adapter in adapters:
                availability = adapter.availability()
                row: dict[str, Any] = {
                    "fixture": entry["id"],
                    "adapter": adapter.name,
                    "media_type": entry["media_type"],
                    "complexity": entry.get("complexity", "unknown"),
                    "source_hash": source_hash,
                    "availability": availability.model_dump(mode="json"),
                    "status": "SKIPPED",
                    "metrics": {},
                }
                if availability.status != "CONFIGURED":
                    row["reason"] = availability.reason or "ADAPTER_NOT_CONFIGURED"
                    rows.append(row)
                    continue
                if not _media_supported(availability, entry["media_type"]):
                    row["reason"] = "MEDIA_NOT_SUPPORTED_BY_ADAPTER"
                    rows.append(row)
                    continue
                request_type = None
                try:
                    from app.ports import DocumentExtractionRequest

                    request_type = DocumentExtractionRequest(
                        source_path=source,
                        source_id=f"fixture:{entry['id']}",
                        source_uri=f"fixture://{entry['id']}",
                        media_type=entry["media_type"],
                        complexity=entry.get("complexity"),
                    )
                    timings: list[float] = []
                    artifact = None
                    for _ in range(max(1, repeats)):
                        started = time.perf_counter()
                        artifact = adapter.extract(request_type)
                        timings.append((time.perf_counter() - started) * 1000)
                    row["status"] = "EXECUTED"
                    row["metrics"] = _metrics(artifact, entry.get("expected", {}))
                    row["metrics"]["processing_time_ms"] = round(sum(timings) / len(timings), 3)
                    row["metrics"]["processing_time_ms_min"] = round(min(timings), 3)
                    row["metrics"]["processing_time_ms_max"] = round(max(timings), 3)
                    row["metrics"]["memory_peak_mb"] = None
                    row["metrics"]["cpu_gpu_requirement"] = "CPU"
                    row["metrics"]["provider_cost"] = 0
                except Exception as error:  # noqa: BLE001  # extraction failures are benchmark evidence
                    row["status"] = "FAILED"
                    row["reason"] = str(error)
                rows.append(row)
    result = {
        "metadata": metadata,
        "status": "EXECUTED",
        "fixtures": rows,
        "selection": {
            "status": "NOT_SELECTED",
            "reason": "benchmark evidence is recorded; no adapter is promoted automatically",
            "configured_selection": False,
        },
    }
    report.parent.mkdir(parents=True, exist_ok=True)
    report.write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--fixtures", type=Path, default=DEFAULT_FIXTURES)
    parser.add_argument("--report", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--repeats", type=int, default=1)
    args = parser.parse_args()
    result = run(args.fixtures.resolve(), args.report.resolve(), max(1, args.repeats))
    print(json.dumps({"status": result["status"], "report": str(args.report.resolve())}, indent=2))
    return 0 if result["status"] in {"EXECUTED", "BLOCKED"} else 1


if __name__ == "__main__":
    raise SystemExit(main())
