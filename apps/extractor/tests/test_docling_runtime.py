from __future__ import annotations

from pathlib import Path

import pytest

from app.adapters.optional import DoclingAdapter
from app.ports import DocumentExtractionRequest


@pytest.mark.provider_runtime
def test_real_docling_provider_preserves_native_html_structure() -> None:
    """Execute the installed Docling provider instead of a synthetic mapper."""

    pytest.importorskip("docling")
    repository_root = Path(__file__).resolve().parents[3]
    source = (
        repository_root
        / "test"
        / "fixtures"
        / "document-intelligence"
        / "complex-layout.html"
    )
    assert source.is_file()

    artifact = DoclingAdapter().extract(
        DocumentExtractionRequest(
            source_path=source,
            source_id="docling-runtime-fixture",
            source_uri="fixture://document-intelligence/complex-layout.html",
            media_type="text/html",
            complexity="complex",
            privacy_policy="LOCAL_ONLY",
        )
    )

    assert artifact.extractor == "docling"
    assert artifact.configuration["mapping"] == "native-docling-document"
    assert artifact.configuration["flattened_before_mapping"] is False
    assert artifact.configuration["native_structure"] is True
    assert artifact.configuration["ocr_requested"] is False
    assert artifact.blocks
    assert artifact.reading_order
    assert all(item_id for item_id in artifact.reading_order)

    text = artifact.text_content()
    assert "Reading order" in text
    assert "Evidence block" in text
    assert "A heading before a table" in text
    assert artifact.tables
    assert any(
        "Signal" in cell
        for table in artifact.tables
        for cell in [*table.headers, *(cell for row in table.rows for cell in row)]
    )

    ids = {item.id for item in artifact.blocks if item.id}
    assert set(artifact.reading_order).issubset(ids)
    assert artifact.quality == "PROVIDER_STRUCTURED"
    assert artifact.quality_metrics["structured_units"] > 0


@pytest.mark.provider_runtime
def test_real_docling_provider_ocr_preserves_scanned_pdf_provenance(
    tmp_path: Path,
) -> None:
    """Force OCR over a raster-only PDF and require grounded native output."""

    pytest.importorskip("docling")
    pytest.importorskip("PIL")
    from PIL import Image, ImageDraw, ImageFont

    source = tmp_path / "scanned-ocr.pdf"
    image = Image.new("RGB", (1654, 2339), "white")
    draw = ImageDraw.Draw(image)
    # Pillow bundles this scalable font on every supported platform.
    font = ImageFont.load_default(size=64)
    draw.multiline_text(
        (120, 260),
        "AKP OCR PROBE 739241\n"
        "Durable document intelligence\n"
        "Native provenance must survive extraction.",
        fill="black",
        font=font,
        spacing=36,
    )
    image.save(source, "PDF", resolution=150.0)

    artifact = DoclingAdapter().extract(
        DocumentExtractionRequest(
            source_path=source,
            source_id="docling-ocr-runtime-fixture",
            source_uri="fixture://document-intelligence/scanned-ocr.pdf",
            media_type="application/pdf",
            complexity="scanned",
            ocr_required=True,
            privacy_policy="LOCAL_ONLY",
            configuration={
                "force_full_page_ocr": True,
                "timeout_seconds": 300,
            },
        )
    )

    normalized = " ".join(artifact.text_content().upper().split())
    assert artifact.extractor == "docling"
    assert artifact.configuration["mapping"] == "native-docling-document"
    assert artifact.configuration["flattened_before_mapping"] is False
    assert artifact.configuration["native_structure"] is True
    assert artifact.configuration["ocr_requested"] is True
    assert artifact.configuration["ocr_engine"] == "provider-default"
    assert "AKP OCR PROBE 739241" in normalized
    assert "DURABLE DOCUMENT INTELLIGENCE" in normalized
    assert artifact.blocks
    assert artifact.pages
    assert any(item.locator.page == 1 for item in artifact.blocks)
    assert any(item.locator.region is not None for item in artifact.blocks)

@pytest.mark.provider_runtime
def test_real_docling_docx_preserves_tables_headings_and_native_locators(
    tmp_path: Path,
) -> None:
    """Native DOCX extraction must not flatten evidence or fabricate page spans."""

    pytest.importorskip("docling")
    from docx import Document

    source = tmp_path / "s1-document-fidelity.docx"
    document = Document()
    document.add_heading("S1 DOCX SOURCE 3157", level=1)
    document.add_paragraph(
        "Preserve structured evidence from a controlled enterprise document."
    )
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "control"
    table.rows[0].cells[1].text = "timeout"
    cells = table.add_row().cells
    cells[0].text = "rollback"
    cells[1].text = "47 minutes"
    document.add_paragraph("End of DOCX evidence 3157")
    document.save(source)

    artifact = DoclingAdapter().extract(
        DocumentExtractionRequest(
            source_path=source,
            source_id="s1-docx-fidelity-3157",
            source_uri="fixture://document-intelligence/s1-docx-fidelity.docx",
            media_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            complexity="table-heavy",
            privacy_policy="LOCAL_ONLY",
        )
    )
    assert artifact.extractor == "docling"
    assert artifact.configuration["mapping"] == "native-docling-document"
    assert artifact.configuration["flattened_before_mapping"] is False
    assert artifact.configuration["native_structure"] is True
    assert artifact.configuration["ocr_requested"] is False
    assert artifact.quality == "PROVIDER_STRUCTURED"

    normalized = " ".join(artifact.text_content().split())
    assert "S1 DOCX SOURCE 3157" in normalized
    assert "controlled enterprise document" in normalized
    assert "End of DOCX evidence 3157" in normalized
    assert artifact.headings
    assert artifact.tables
    assert any(
        "47 minutes" in cell
        for table in artifact.tables
        for row in table.rows
        for cell in row
    )
    assert artifact.reading_order
    identifiers = {item.id for item in artifact.blocks if item.id}
    assert set(artifact.reading_order).issubset(identifiers)
    assert all(item.locator.source_hash for item in artifact.blocks)
    assert artifact.quality_metrics["structured_units"] > 0

@pytest.mark.provider_runtime
def test_real_docling_digital_pdf_preserves_visible_text_and_page_locator(
    tmp_path: Path,
) -> None:
    """Digital PDF provenance is independent from the forced-OCR pipeline."""

    pytest.importorskip("docling")
    source = tmp_path / "s1-digital-fidelity.pdf"
    stream = b"BT\n/F1 24 Tf\n72 600 Td\n(S1 PDF PROVENANCE 3179) Tj\nET\n"
    second_page = b"BT\n/F1 24 Tf\n72 600 Td\n(S1 SECOND PAGE 4281) Tj\nET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R 6 0 R] /Count 2 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length "
        + str(len(stream)).encode("ascii")
        + b" >>\nstream\n"
        + stream
        + b"endstream",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 7 0 R >>"
        ),
        b"<< /Length "
        + str(len(second_page)).encode("ascii")
        + b" >>\nstream\n"
        + second_page
        + b"endstream",
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, item in enumerate(objects, start=1):
        offsets.append(len(output))
        output.extend(
            str(index).encode("ascii") + b" 0 obj\n" + item + b"\nendobj\n"
        )
    xref = len(output)
    output.extend(b"xref\n0 8\n0000000000 65535 f \n")
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    output.extend(
        b"trailer\n<< /Size 8 /Root 1 0 R >>\nstartxref\n"
        + str(xref).encode("ascii")
        + b"\n%%EOF\n"
    )
    source.write_bytes(output)

    artifact = DoclingAdapter().extract(
        DocumentExtractionRequest(
            source_path=source,
            source_id="s1-digital-pdf-3179",
            source_uri="fixture://document-intelligence/s1-digital-fidelity.pdf",
            media_type="application/pdf",
            complexity="simple",
            privacy_policy="LOCAL_ONLY",
        )
    )
    assert artifact.extractor == "docling"
    assert artifact.configuration["mapping"] == "native-docling-document"
    assert artifact.configuration["flattened_before_mapping"] is False
    assert artifact.configuration["ocr_requested"] is False
    assert artifact.quality == "PROVIDER_STRUCTURED"
    text = " ".join(artifact.text_content().split())
    assert "S1 PDF PROVENANCE 3179" in text
    assert "S1 SECOND PAGE 4281" in text
    assert artifact.pages
    assert artifact.blocks
    assert artifact.reading_order
    assert any(item.locator.page == 1 for item in artifact.blocks)
    assert any(item.locator.page == 2 for item in artifact.blocks)
    assert all(item.locator.source_hash for item in artifact.blocks)

