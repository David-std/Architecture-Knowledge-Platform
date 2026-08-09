from pathlib import Path

import pytest
from docx import Document
from fastapi import HTTPException
from pptx import Presentation

from app.main import capabilities, extract
from app.models import ExtractRequest


def test_capabilities_are_explicit() -> None:
    by_media = {
        str(item["media"]): item
        for item in capabilities()["capabilities"]  # type: ignore[index]
    }
    assert by_media["docx"]["status"] == "CONFIGURED"
    assert by_media["pptx"]["status"] == "CONFIGURED"
    assert by_media["audio-transcript"]["status"] == "CAPABILITY_NOT_CONFIGURED"
    assert by_media["video-transcript-visual"]["timestamps"] is False


def test_docx_and_pptx_extract_with_locators(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("AKP_EXTRACTOR_TOKEN", "expected")
    monkeypatch.setenv("AKP_EXTRACTOR_ROOTS", str(tmp_path))

    docx_path = tmp_path / "sample.docx"
    document = Document()
    document.add_paragraph("Invariant evidence")
    document.save(docx_path)
    docx = extract(ExtractRequest(source_uri=str(docx_path)), "expected")
    assert docx.extractor == "python-docx"
    assert docx.artifacts[0].content == "Invariant evidence"
    assert docx.artifacts[0].locator["paragraph"] == 1

    pptx_path = tmp_path / "sample.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Architecture decision"
    presentation.save(pptx_path)
    pptx = extract(ExtractRequest(source_uri=str(pptx_path)), "expected")
    assert pptx.extractor == "python-pptx"
    assert "Architecture decision" in (pptx.artifacts[0].content or "")
    assert pptx.artifacts[0].locator["slide"] == 1


def test_audio_returns_capability_not_configured(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("AKP_EXTRACTOR_TOKEN", "expected")
    monkeypatch.setenv("AKP_EXTRACTOR_ROOTS", str(tmp_path))
    source = tmp_path / "audio.mp3"
    source.write_bytes(b"not-real-audio")
    with pytest.raises(HTTPException) as error:
        extract(ExtractRequest(source_uri=str(source), media_type="audio/mpeg"), "expected")
    assert error.value.status_code == 501
    assert error.value.detail["code"] == "CAPABILITY_NOT_CONFIGURED"
