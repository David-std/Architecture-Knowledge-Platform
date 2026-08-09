from hashlib import sha256
from pathlib import Path

from docx import Document

from ..models import Artifact


def extract_docx(path: Path) -> list[Artifact]:
    source_hash = sha256(path.read_bytes()).hexdigest()
    document = Document(str(path))
    artifacts: list[Artifact] = []
    for index, paragraph in enumerate(document.paragraphs, start=1):
        if not paragraph.text.strip():
            continue
        artifacts.append(
            Artifact(
                kind="docx-paragraph-text",
                content=paragraph.text,
                locator={
                    "kind": "docx",
                    "paragraph": index,
                    "source_hash": source_hash,
                },
                quality="EXACT_TEXT",
            )
        )
    for table_index, table in enumerate(document.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            text = " | ".join(cell.text.strip() for cell in row.cells)
            if not text.strip(" |"):
                continue
            artifacts.append(
                Artifact(
                    kind="docx-table-row-text",
                    content=text,
                    locator={
                        "kind": "docx",
                        "table": table_index,
                        "row": row_index,
                        "source_hash": source_hash,
                    },
                    quality="EXACT_TEXT",
                )
            )
    if not artifacts:
        artifacts.append(
            Artifact(
                kind="docx-empty",
                locator={"kind": "docx", "source_hash": source_hash},
                warnings=["NO_TEXT_EXTRACTED"],
                quality="MACHINE_EXTRACTED",
            )
        )
    return artifacts
