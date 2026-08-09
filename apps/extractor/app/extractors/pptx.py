from hashlib import sha256
from pathlib import Path

from pptx import Presentation

from ..models import Artifact


def extract_pptx(path: Path) -> list[Artifact]:
    source_hash = sha256(path.read_bytes()).hexdigest()
    presentation = Presentation(str(path))
    artifacts: list[Artifact] = []
    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = [
            str(shape.text).strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and str(shape.text).strip()
        ]
        notes = ""
        if slide.has_notes_slide:
            notes = "\n".join(
                str(shape.text).strip()
                for shape in slide.notes_slide.notes_text_frame.paragraphs
                if str(shape.text).strip()
            )
        content = "\n".join([*text_parts, notes]).strip()
        artifacts.append(
            Artifact(
                kind="pptx-slide-text",
                content=content or None,
                locator={
                    "kind": "pptx",
                    "slide": slide_number,
                    "source_hash": source_hash,
                },
                warnings=[] if content else ["NO_TEXT_EXTRACTED"],
                quality="EXACT_TEXT" if content else "MACHINE_EXTRACTED",
            )
        )
    return artifacts
