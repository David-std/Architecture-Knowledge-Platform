"""Deterministic local document parser.

This adapter intentionally favours reproducibility and source locators over
aggressive layout inference.  It is a useful baseline for every supported
media type and remains functional when optional document-intelligence
providers are not installed.
"""

from __future__ import annotations

import re
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass, field
from typing import Any, ClassVar

from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from pptx import Presentation
from pypdf import PdfReader

from ..models import (
    ArtifactItem,
    BoundingBox,
    DocumentArtifact,
    PageArtifact,
    StructuralLocator,
    TableArtifact,
)
from ..ports import (
    AdapterAvailability,
    CapabilityStatus,
    DocumentExtractionRequest,
    DocumentIntelligencePort,
    UnsupportedMediaType,
)
from .base import infer_media_type, sha256_path

_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_LIST = re.compile(r"^\s*(?:[-*+]\s+|\d+[.)]\s+)(.+?)\s*$")
_IMAGE = re.compile(r"!\[([^]]*)\]\(([^)]+)\)")
_EQUATION = re.compile(r"^(?:\${1,2}|\\\[|\\\(|\\begin\{).*(?:\${1,2}|\\\]|\\\)|\\end\{.*\})?$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{3,}:?\s*(?:\|\s*:?-{3,}:?\s*)+\|?\s*$")


@dataclass
class _ParsedUnits:
    blocks: list[ArtifactItem | TableArtifact] = field(default_factory=list)
    headings: list[ArtifactItem] = field(default_factory=list)
    paragraphs: list[ArtifactItem] = field(default_factory=list)
    lists: list[ArtifactItem] = field(default_factory=list)
    tables: list[TableArtifact] = field(default_factory=list)
    figures: list[ArtifactItem] = field(default_factory=list)
    equations: list[ArtifactItem] = field(default_factory=list)
    code: list[ArtifactItem] = field(default_factory=list)
    reading_order: list[str] = field(default_factory=list)
    locators: list[StructuralLocator] = field(default_factory=list)


def _split_table_row(line: str) -> list[str]:
    value = line.strip()
    value = value.removeprefix("|").removesuffix("|")
    return [cell.strip() for cell in value.split("|")]


def _new_locator(
    source_hash: str,
    source_ref: str,
    *,
    kind: str,
    line_start: int | None = None,
    line_end: int | None = None,
    page: int | None = None,
    slide: int | None = None,
    paragraph: int | None = None,
    table: int | None = None,
    row: int | None = None,
    sheet: str | None = None,
    heading_path: list[str] | None = None,
    region: dict[str, Any] | None = None,
) -> StructuralLocator:
    values: dict[str, Any] = {
        "kind": kind,
        "source_hash": source_hash,
        "path": source_ref,
        "start_line": line_start,
        "end_line": line_end,
        "page": page,
        "slide": slide,
        "paragraph": paragraph,
        "table": table,
        "row": row,
        "sheet": sheet,
        "heading_path": heading_path or [],
    }
    if region:
        values["region"] = region
    return StructuralLocator.model_validate({key: value for key, value in values.items() if value is not None})


def _parse_text_units(
    content: str,
    *,
    source_hash: str,
    source_ref: str,
    id_prefix: str,
    page: int | None = None,
    slide: int | None = None,
    line_offset: int = 0,
) -> _ParsedUnits:
    result = _ParsedUnits()
    lines = content.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    heading_stack: list[str] = []
    sequence = 0

    def make_id(kind: str, line: int) -> str:
        nonlocal sequence
        sequence += 1
        return f"{id_prefix}-{kind}-{line}-{sequence}"

    def append_item(item: ArtifactItem | TableArtifact, category: list[Any]) -> None:
        result.blocks.append(item)
        category.append(item)
        if item.id:
            result.reading_order.append(item.id)
        result.locators.append(item.locator)

    pending: list[tuple[int, str]] = []
    in_code = False
    code_start = 0
    code_language = ""
    code_lines: list[str] = []

    def flush_paragraph() -> None:
        if not pending:
            return
        first_line = pending[0][0]
        last_line = pending[-1][0]
        text = "\n".join(value for _, value in pending).strip()
        pending.clear()
        if not text:
            return
        locator = _new_locator(
            source_hash,
            source_ref,
            kind="paragraph",
            line_start=first_line + line_offset,
            line_end=last_line + line_offset,
            page=page,
            slide=slide,
            heading_path=heading_stack,
        )
        item = ArtifactItem(
            id=make_id("paragraph", first_line),
            kind="paragraph",
            text=text,
            locator=locator,
            metadata={"heading_path": [*heading_stack]},
        )
        append_item(item, result.paragraphs)

    line_number = 0
    while line_number < len(lines):
        raw_line = lines[line_number]
        line = raw_line.strip()
        current = line_number + 1

        if line.startswith(("```", "~~~")):
            flush_paragraph()
            marker = line[:3]
            if not in_code:
                in_code = True
                code_start = current
                code_language = line[3:].strip()
                code_lines = []
            elif line.startswith(marker):
                locator = _new_locator(
                    source_hash,
                    source_ref,
                    kind="code",
                    line_start=code_start + line_offset,
                    line_end=current + line_offset,
                    page=page,
                    slide=slide,
                    heading_path=heading_stack,
                )
                item = ArtifactItem(
                    id=make_id("code", code_start),
                    kind="code",
                    text="\n".join(code_lines),
                    locator=locator,
                    metadata={"language": code_language, "heading_path": [*heading_stack]},
                )
                append_item(item, result.code)
                in_code = False
                code_lines = []
            else:
                code_lines.append(raw_line)
            line_number += 1
            continue
        if in_code:
            code_lines.append(raw_line)
            line_number += 1
            continue

        heading = _HEADING.match(line)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            title = heading.group(2).strip()
            heading_stack = heading_stack[: level - 1] + [title]
            locator = _new_locator(
                source_hash,
                source_ref,
                kind="heading",
                line_start=current + line_offset,
                line_end=current + line_offset,
                page=page,
                slide=slide,
                heading_path=heading_stack,
            )
            item = ArtifactItem(
                id=make_id("heading", current),
                kind="heading",
                text=title,
                locator=locator,
                metadata={"level": level, "heading_path": [*heading_stack]},
            )
            append_item(item, result.headings)
            line_number += 1
            continue

        if line and line_number + 1 < len(lines) and "|" in line and _TABLE_SEPARATOR.match(lines[line_number + 1]):
            flush_paragraph()
            table_start = current
            table_lines = [line]
            line_number += 2
            while line_number < len(lines) and "|" in lines[line_number] and lines[line_number].strip():
                table_lines.append(lines[line_number].strip())
                line_number += 1
            headers = _split_table_row(table_lines[0])
            rows = [_split_table_row(row) for row in table_lines[1:]]
            locator = _new_locator(
                source_hash,
                source_ref,
                kind="table",
                line_start=table_start + line_offset,
                line_end=table_start + len(table_lines) + 1 + line_offset,
                page=page,
                slide=slide,
                heading_path=heading_stack,
            )
            table = TableArtifact(
                id=make_id("table", table_start),
                kind="table",
                text=" | ".join(headers),
                locator=locator,
                headers=headers,
                rows=rows,
                metadata={"heading_path": [*heading_stack]},
            )
            append_item(table, result.tables)
            continue

        list_match = _LIST.match(line)
        if list_match:
            flush_paragraph()
            locator = _new_locator(
                source_hash,
                source_ref,
                kind="list-item",
                line_start=current + line_offset,
                line_end=current + line_offset,
                page=page,
                slide=slide,
                heading_path=heading_stack,
            )
            item = ArtifactItem(
                id=make_id("list", current),
                kind="list-item",
                text=list_match.group(1),
                locator=locator,
                metadata={"heading_path": [*heading_stack]},
            )
            append_item(item, result.lists)
            line_number += 1
            continue

        image_match = _IMAGE.search(line)
        equation = bool(_EQUATION.match(line))
        if image_match:
            flush_paragraph()
            locator = _new_locator(
                source_hash,
                source_ref,
                kind="figure",
                line_start=current + line_offset,
                line_end=current + line_offset,
                page=page,
                slide=slide,
                heading_path=heading_stack,
            )
            item = ArtifactItem(
                id=make_id("figure", current),
                kind="figure",
                text=image_match.group(1) or None,
                locator=locator,
                metadata={"uri": image_match.group(2), "heading_path": [*heading_stack]},
            )
            append_item(item, result.figures)
            line_number += 1
            continue

        if equation:
            flush_paragraph()
            locator = _new_locator(
                source_hash,
                source_ref,
                kind="equation",
                line_start=current + line_offset,
                line_end=current + line_offset,
                page=page,
                slide=slide,
                heading_path=heading_stack,
            )
            item = ArtifactItem(
                id=make_id("equation", current),
                kind="equation",
                text=line,
                locator=locator,
                metadata={"heading_path": [*heading_stack]},
            )
            append_item(item, result.equations)
            line_number += 1
            continue

        if not line:
            flush_paragraph()
        else:
            pending.append((current, raw_line.rstrip()))
        line_number += 1
    if in_code:
        locator = _new_locator(
            source_hash,
            source_ref,
            kind="code",
            line_start=code_start + line_offset,
            line_end=len(lines) + line_offset,
            page=page,
            slide=slide,
            heading_path=heading_stack,
        )
        item = ArtifactItem(
            id=make_id("code", code_start),
            kind="code",
            text="\n".join(code_lines),
            locator=locator,
            metadata={"language": code_language, "unterminated": True},
        )
        append_item(item, result.code)
    flush_paragraph()
    return result


def _artifact_from_units(
    request: DocumentExtractionRequest,
    source_hash: str,
    *,
    extractor: str,
    version: str,
    media_type: str,
    units: Iterable[_ParsedUnits],
    pages: list[PageArtifact] | None = None,
    warnings: list[str] | None = None,
    quality: str = "EXACT_TEXT",
) -> DocumentArtifact:
    merged = _ParsedUnits()
    for unit in units:
        merged.blocks.extend(unit.blocks)
        merged.headings.extend(unit.headings)
        merged.paragraphs.extend(unit.paragraphs)
        merged.lists.extend(unit.lists)
        merged.tables.extend(unit.tables)
        merged.figures.extend(unit.figures)
        merged.equations.extend(unit.equations)
        merged.code.extend(unit.code)
        merged.reading_order.extend(unit.reading_order)
        merged.locators.extend(unit.locators)
    source_locator = _new_locator(source_hash, request.source_uri or request.source_id, kind="source")
    merged.locators.insert(0, source_locator)
    all_pages = pages or []
    merged.locators.extend(page.locator for page in all_pages)
    if not merged.blocks and not all_pages:
        warnings = [*(warnings or []), "NO_TEXT_EXTRACTED"]
    return DocumentArtifact(
        source_id=request.source_id,
        source_hash=source_hash,
        media_type=media_type,
        extractor=extractor,
        extractor_version=version,
        configuration={"adapter": "deterministic-baseline", **request.configuration},
        pages=all_pages,
        blocks=merged.blocks,
        headings=merged.headings,
        paragraphs=merged.paragraphs,
        lists=merged.lists,
        tables=merged.tables,
        figures=merged.figures,
        equations=merged.equations,
        code=merged.code,
        reading_order=merged.reading_order,
        locators=merged.locators,
        warnings=warnings or [],
        quality=quality,
        quality_metrics={
            "structured_units": len(merged.blocks),
            "locator_count": len(merged.locators),
            "pages": len(all_pages),
        },
    )


def artifact_from_text(
    request: DocumentExtractionRequest,
    content: str,
    *,
    extractor: str = "deterministic-text",
    media_type: str | None = None,
    source_hash: str | None = None,
    warnings: list[str] | None = None,
) -> DocumentArtifact:
    digest = source_hash or sha256_path(request.source_path)
    media = media_type or infer_media_type(request.source_path, request.media_type)
    units = _parse_text_units(
        content,
        source_hash=digest,
        source_ref=request.source_uri or request.source_id,
        id_prefix="document",
    )
    page_locator = _new_locator(digest, request.source_uri or request.source_id, kind="page", page=1)
    page = PageArtifact(id="page-1", kind="page", page=1, text=content, locator=page_locator)
    return _artifact_from_units(
        request,
        digest,
        extractor=extractor,
        version="0.3.0",
        media_type=media,
        units=[units],
        pages=[page],
        warnings=warnings,
    )


class DeterministicTextAdapter(DocumentIntelligencePort):
    name = "deterministic-baseline"
    version = "0.3.0"
    media_suffixes: ClassVar[set[str]] = {
        ".md",
        ".markdown",
        ".txt",
        ".csv",
        ".json",
        ".yaml",
        ".yml",
        ".xml",
        ".py",
        ".ts",
        ".tsx",
        ".js",
        ".jsx",
        ".java",
        ".cs",
        ".sql",
        ".html",
        ".htm",
        ".pdf",
        ".docx",
        ".pptx",
        ".xlsx",
        ".png",
        ".jpg",
        ".jpeg",
        ".webp",
        ".gif",
    }

    def availability(self) -> AdapterAvailability:
        return AdapterAvailability(
            adapter=self.name,
            version=self.version,
            status=CapabilityStatus.CONFIGURED,
            reason="stdlib-and-declared-local-parsers",
            media=[
                "markdown-text",
                "code",
                "html-snapshot",
                "pdf",
                "docx",
                "pptx",
                "xlsx",
                "image-metadata",
            ],
            complexities=["simple", "digital", "complex", "scanned", "formula", "table-heavy"],
            locators=True,
            structured_output=True,
            local=True,
        )

    def extract(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        path = request.source_path
        suffix = path.suffix.lower()
        media = infer_media_type(path, request.media_type)
        if suffix in {".html", ".htm"} or media == "text/html":
            return self._extract_html(request)
        if (
            suffix in self.media_suffixes
            and suffix
            not in {
                ".html",
                ".htm",
                ".pdf",
                ".docx",
                ".pptx",
                ".xlsx",
                ".png",
                ".jpg",
                ".jpeg",
                ".webp",
                ".gif",
            }
        ) or media.startswith("text/") or media in {
            "application/json",
            "application/xml",
            "application/javascript",
            "application/x-yaml",
            "application/yaml",
        }:
            return artifact_from_text(
                request,
                path.read_text(encoding="utf-8", errors="replace"),
                extractor="deterministic-text",
                media_type=media,
            )
        if suffix == ".pdf" or media == "application/pdf":
            return self._extract_pdf(request)
        if suffix == ".docx":
            return self._extract_docx(request)
        if suffix == ".pptx":
            return self._extract_pptx(request)
        if suffix == ".xlsx":
            return self._extract_xlsx(request)
        if suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif"} or media.startswith("image/"):
            return self._extract_image(request)
        raise UnsupportedMediaType(f"unsupported media type: {media}")

    def _extract_html(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        raw = request.source_path.read_bytes()
        soup = BeautifulSoup(raw, "html.parser")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        lines: list[str] = []
        for element in soup.find_all(["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre"]):
            text = element.get_text(" ", strip=True)
            if not text:
                continue
            if element.name and element.name.startswith("h"):
                lines.append(f"{'#' * int(element.name[1:])} {text}")
            elif element.name == "li":
                lines.append(f"- {text}")
            else:
                lines.append(text)
        if not lines:
            lines = [line.strip() for line in soup.get_text("\n").splitlines() if line.strip()]
        return artifact_from_text(
            request,
            "\n\n".join(lines),
            extractor="deterministic-html",
            media_type="text/html",
            source_hash=sha256_path(request.source_path),
        )

    def _extract_pdf(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        digest = sha256_path(request.source_path)
        reader = PdfReader(str(request.source_path))
        units: list[_ParsedUnits] = []
        pages: list[PageArtifact] = []
        warnings: list[str] = []
        for page_number, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if not text.strip():
                warnings.append(f"PAGE_{page_number}_NO_TEXT_EXTRACTED")
            page_locator = _new_locator(
                digest,
                request.source_uri or request.source_id,
                kind="page",
                page=page_number,
            )
            pages.append(
                PageArtifact(
                    id=f"page-{page_number}",
                    kind="page",
                    page=page_number,
                    text=text,
                    locator=page_locator,
                )
            )
            units.append(
                _parse_text_units(
                    text,
                    source_hash=digest,
                    source_ref=request.source_uri or request.source_id,
                    id_prefix=f"page-{page_number}",
                    page=page_number,
                )
            )
        if any(not (page.text or "").strip() for page in pages):
            warnings.append("LAYOUT_AND_OCR_LIMITED_TO_DETERMINISTIC_BASELINE")
        return _artifact_from_units(
            request,
            digest,
            extractor="deterministic-pypdf",
            version=self.version,
            media_type="application/pdf",
            units=units,
            pages=pages,
            warnings=warnings,
            quality="MACHINE_EXTRACTED",
        )

    def _extract_docx(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        digest = sha256_path(request.source_path)
        document = Document(str(request.source_path))
        units = _ParsedUnits()
        heading_stack: list[str] = []
        for paragraph_number, paragraph in enumerate(document.paragraphs, start=1):
            text = paragraph.text.strip()
            if not text:
                continue
            style = paragraph.style.name if paragraph.style else ""
            heading_match = re.search(r"heading\s*(\d+)", style, re.IGNORECASE)
            if heading_match:
                level = int(heading_match.group(1))
                heading_stack = heading_stack[: level - 1] + [text]
                kind = "heading"
                metadata = {"level": level, "style": style, "heading_path": [*heading_stack]}
            else:
                kind = "paragraph"
                metadata = {"style": style, "heading_path": [*heading_stack]}
            locator = _new_locator(
                digest,
                request.source_uri or request.source_id,
                kind="docx-paragraph",
                paragraph=paragraph_number,
                heading_path=heading_stack,
            )
            item = ArtifactItem(
                id=f"paragraph-{paragraph_number}",
                kind=kind,
                text=text,
                locator=locator,
                metadata=metadata,
            )
            units.blocks.append(item)
            units.locators.append(locator)
            units.reading_order.append(item.id or "")
            (units.headings if kind == "heading" else units.paragraphs).append(item)
        for table_number, table in enumerate(document.tables, start=1):
            rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
            if not rows:
                continue
            locator = _new_locator(
                digest,
                request.source_uri or request.source_id,
                kind="docx-table",
                table=table_number,
                row=1,
            )
            table_item = TableArtifact(
                id=f"table-{table_number}",
                kind="table",
                text=" | ".join(rows[0]),
                locator=locator,
                headers=rows[0],
                rows=rows[1:],
            )
            units.blocks.append(table_item)
            units.tables.append(table_item)
            units.locators.append(locator)
            units.reading_order.append(table_item.id or "")
        return _artifact_from_units(
            request,
            digest,
            extractor="python-docx",
            version=self.version,
            media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            units=[units],
            quality="EXACT_TEXT",
        )

    def _extract_pptx(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        digest = sha256_path(request.source_path)
        presentation = Presentation(str(request.source_path))
        units: list[_ParsedUnits] = []
        pages: list[PageArtifact] = []
        for slide_number, slide in enumerate(presentation.slides, start=1):
            text_parts = [
                str(shape.text).strip()
                for shape in slide.shapes
                if hasattr(shape, "text") and str(shape.text).strip()
            ]
            if slide.has_notes_slide:
                text_parts.extend(
                    str(paragraph.text).strip()
                    for paragraph in slide.notes_slide.notes_text_frame.paragraphs
                    if str(paragraph.text).strip()
                )
            text = "\n".join(text_parts).strip()
            locator = _new_locator(
                digest,
                request.source_uri or request.source_id,
                kind="pptx-slide",
                slide=slide_number,
            )
            pages.append(
                PageArtifact(
                    id=f"slide-{slide_number}",
                    kind="slide",
                    page=slide_number,
                    text=text,
                    locator=locator,
                )
            )
            units.append(
                _parse_text_units(
                    text,
                    source_hash=digest,
                    source_ref=request.source_uri or request.source_id,
                    id_prefix=f"slide-{slide_number}",
                    slide=slide_number,
                )
            )
        return _artifact_from_units(
            request,
            digest,
            extractor="python-pptx",
            version=self.version,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            units=units,
            pages=pages,
            warnings=["SLIDE_SHAPE_GEOMETRY_NOT_EXTRACTED"],
            quality="MACHINE_EXTRACTED",
        )

    def _extract_image(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        digest = sha256_path(request.source_path)
        with Image.open(request.source_path) as image:
            width, height, image_format, mode = image.width, image.height, image.format, image.mode
        locator = _new_locator(
            digest,
            request.source_uri or request.source_id,
            kind="image",
            region={"x": 0, "y": 0, "width": width, "height": height, "unit": "pixel"},
        )
        figure = ArtifactItem(
            id="figure-1",
            kind="figure",
            text=None,
            locator=locator,
            metadata={"width": width, "height": height, "format": image_format, "mode": mode},
        )
        return DocumentArtifact(
            source_id=request.source_id,
            source_hash=digest,
            media_type=infer_media_type(request.source_path, request.media_type),
            extractor="pillow-metadata",
            extractor_version=self.version,
            configuration={"adapter": "deterministic-baseline", **request.configuration},
            figures=[figure],
            bounding_boxes=[BoundingBox(x=0, y=0, width=width, height=height, unit="pixel")],
            reading_order=["figure-1"],
            locators=[locator],
            warnings=["OCR_NOT_CONFIGURED"],
            quality="DETERMINISTIC_METADATA",
            quality_metrics={"width": width, "height": height},
        )

    def _extract_xlsx(self, request: DocumentExtractionRequest) -> DocumentArtifact:
        digest = sha256_path(request.source_path)
        shared_strings: list[str] = []
        tables: list[TableArtifact] = []
        blocks: list[ArtifactItem | TableArtifact] = []
        locators: list[StructuralLocator] = []
        reading_order: list[str] = []
        with zipfile.ZipFile(request.source_path) as archive:
            if "xl/sharedStrings.xml" in archive.namelist():
                root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
                namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for item in root.findall("x:si", namespace):
                    shared_strings.append("".join(item.itertext()).strip())
            workbook = ET.fromstring(archive.read("xl/workbook.xml"))
            namespace = {"x": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
            sheets = workbook.findall("x:sheets/x:sheet", namespace)
            for sheet_number, sheet in enumerate(sheets, start=1):
                sheet_name = sheet.attrib.get("name", f"Sheet{sheet_number}")
                worksheet_name = f"xl/worksheets/sheet{sheet_number}.xml"
                if worksheet_name not in archive.namelist():
                    continue
                rows: dict[int, dict[int, str]] = {}
                root = ET.fromstring(archive.read(worksheet_name))
                for row in root.findall(".//x:sheetData/x:row", namespace):
                    row_number = int(row.attrib.get("r", "1"))
                    values: dict[int, str] = {}
                    for cell in row.findall("x:c", namespace):
                        reference = cell.attrib.get("r", "A1")
                        column_match = re.match(r"[A-Za-z]+", reference)
                        if column_match is None:
                            continue
                        column = 0
                        for character in column_match.group(0):
                            column = column * 26 + ord(character.upper()) - 64
                        value = cell.find("x:v", namespace)
                        inline = cell.find("x:is", namespace)
                        text = "" if value is None else "".join(value.itertext()).strip()
                        if inline is not None:
                            text = "".join(inline.itertext()).strip()
                        if cell.attrib.get("t") == "s" and text:
                            text = shared_strings[int(text)]
                        formula = cell.find("x:f", namespace)
                        if formula is not None and formula.text:
                            text = f"={formula.text}" + (f" ({text})" if text else "")
                        values[column] = text
                    if values:
                        rows[row_number] = values
                if not rows:
                    continue
                max_column = max(max(values) for values in rows.values())
                matrix = [
                    [rows[row_number].get(column, "") for column in range(1, max_column + 1)]
                    for row_number in sorted(rows)
                ]
                headers = matrix[0]
                data_rows = matrix[1:]
                locator = _new_locator(
                    digest,
                    request.source_uri or request.source_id,
                    kind="xlsx-table",
                    sheet=sheet_name,
                    row=1,
                )
                table = TableArtifact(
                    id=f"table-{sheet_number}",
                    kind="table",
                    text=" | ".join(headers),
                    locator=locator,
                    headers=headers,
                    rows=data_rows,
                    metadata={"sheet": sheet_name, "sheet_index": sheet_number},
                )
                tables.append(table)
                blocks.append(table)
                locators.append(locator)
                reading_order.append(table.id or "")
        return DocumentArtifact(
            source_id=request.source_id,
            source_hash=digest,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            extractor="stdlib-xlsx",
            extractor_version=self.version,
            configuration={"adapter": "deterministic-baseline", **request.configuration},
            blocks=blocks,
            tables=tables,
            reading_order=reading_order,
            locators=locators,
            warnings=[] if tables else ["NO_WORKSHEET_VALUES"],
            quality="STRUCTURED_TABLE",
            quality_metrics={"sheets": len(tables), "tables": len(tables)},
        )
